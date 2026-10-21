"""Layout QA for a rendered deck: finds text that runs off the page or past the footer rule, out of its tinted
box, onto a screenshot, or into other text. Text positions come from the PDF (what people actually see); box and
picture positions come from the PPTX (LibreOffice rasterizes shapes in the PDF, so they can't be told apart there).

    python renderer/deck_qa.py <deck.pptx> [...]     (the PDF next to it is used; exit 1 if anything is found)
"""
import sys
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

LIGHT = 'F5F3FF'
FOOT_Y = 7.02 * 72          # the footer rule, PDF points (slides are 13.333 x 7.5 in = 960 x 540 pt)
EPS = 1.5
EMU_PT = 72 / 914400


def shape_rects(slide):
    boxes, pics = [], []
    for sh in slide.shapes:
        r = pymupdf.Rect(sh.left * EMU_PT, sh.top * EMU_PT, (sh.left + sh.width) * EMU_PT, (sh.top + sh.height) * EMU_PT)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and r.width > 100:
            pics.append(r)
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and r.width > 72:
            try:
                if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == LIGHT:
                    boxes.append(r)
            except (AttributeError, TypeError):
                pass
    return boxes, pics


def text_lines(page):
    out = []
    for b in page.get_text('dict')['blocks']:
        for ln in b.get('lines', []):
            text = ''.join(s['text'] for s in ln['spans']).strip()
            if text:
                out.append((pymupdf.Rect(ln['bbox']), text, max(s['size'] for s in ln['spans'])))
    return out


def check_page(page, slide):
    issues = []
    boxes, pics = shape_rects(slide)
    lines = text_lines(page)
    for r, text, size in lines:
        footer = r.y0 >= FOOT_Y - 1 and size <= 9.5
        if r.y1 > page.rect.height - EPS or (not footer and r.y1 > FOOT_Y):
            issues.append(f'runs off the bottom: "{text[:70]}"')
        for bx in boxes:
            if bx.x0 - EPS <= r.x0 and r.x1 <= bx.x1 + EPS and bx.y0 - EPS <= r.y0 <= bx.y1 and r.y1 > bx.y1 + EPS:
                issues.append(f'runs out of its box: "{text[:70]}"')
        for p in pics:
            if (r & p).get_area() > 0.2 * r.get_area():
                issues.append(f'runs onto a screenshot: "{text[:70]}"')
    for i, (a, ta, _) in enumerate(lines):
        for b, tb, _ in lines[i + 1:]:
            if (a & b).get_area() > 0.3 * min(a.get_area(), b.get_area()):
                issues.append(f'text overlaps text: "{ta[:40]}" / "{tb[:40]}"')
    return issues


def check(pptx_path):
    pptx_path = Path(pptx_path)
    prs = Presentation(str(pptx_path))
    doc = pymupdf.open(str(pptx_path.with_suffix('.pdf')))
    found = []
    for i in range(min(len(prs.slides), doc.page_count)):
        found += [(i + 1, m) for m in check_page(doc[i], prs.slides[i])]
    return found, doc.page_count


def main(paths):
    bad = 0
    for path in paths:
        found, n = check(path)
        for pg, msg in found:
            print(f'{Path(path).name} p{pg}: {msg}')
        print(f'== {Path(path).name}: {n} pages, {len(found)} issue(s)')
        bad += len(found)
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
